"""gov.cn article extraction and catalog update decisions."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import official_fulltext as fulltext

PAGE = """
<html><body>
<div id="UCAP-CONTENT">
<p>国家药监局关于发布关节镜下无源手术器械产品等8项注册审查指导原则的通告（2021年第102号）</p>
<p>特此通告。</p>
<p>附件：<a href="./P020210001.doc">指导原则</a></p>
</div>
<a href="/unrelated.docx">页脚</a>
</body></html>
"""


def test_titles_and_section():
    notice = "国家药监局关于发布医疗器械出口销售证明管理规定的公告"
    assert fulltext.document_titles(notice) == ["医疗器械出口销售证明管理规定"]
    assert fulltext.target_section("医疗器械出口销售证明管理规定") == "nmpa/regulations"
    many = "关于发布《关节镜下无源手术器械产品注册审查指导原则》《麻醉面罩产品注册审查指导原则》的通告"
    assert fulltext.document_titles(many) == [
        "关节镜下无源手术器械产品注册审查指导原则",
        "麻醉面罩产品注册审查指导原则",
    ]
    assert fulltext.target_section("关节镜下无源手术器械产品注册审查指导原则") == "nmpa/guidance"


def test_parse_article_keeps_content_attachment_only():
    text, links = fulltext.parse_article(
        PAGE, "https://www.gov.cn/zhengce/zhengceku/2021-12/17/content_5661507.htm"
    )
    assert "特此通告" in text
    assert links == ["https://www.gov.cn/zhengce/zhengceku/2021-12/17/P020210001.doc"]
    assert fulltext.document_number(text) == "2021年第102号"


def test_plan_keeps_long_page_and_writes_missing_fulltext():
    long_page = {"body_len": 4000, "source_url": "https://reguverse.com/documentation/fsc/"}
    assert fulltext.plan_document(long_page, "甲" * 8000) == "retarget"
    official = {"body_len": 4000, "source_url": "https://www.gov.cn/zhengce/content.htm"}
    assert fulltext.plan_document(official, "甲" * 8000) == "present"
    assert fulltext.plan_document(None, "甲" * 800) == "create"
    assert fulltext.plan_document(None, "只有一段通告。") == "no_fulltext"
    short = {"body_len": 80, "source_url": ""}
    assert fulltext.plan_document(short, "甲" * 900) == "replace"


def test_split_numbered_attachments():
    body = (
        "附件1\n\n关节镜下无源手术器械产品\n\n注册审查指导原则\n\n"
        + ("甲" * 700)
        + "\n附件2\n\n麻醉面罩产品注册审查指导原则\n\n"
        + ("乙" * 700)
    )
    pieces = fulltext.split_attachments(body)
    assert [title for title, _ in pieces] == [
        "关节镜下无源手术器械产品注册审查指导原则",
        "麻醉面罩产品注册审查指导原则",
    ]


def test_split_named_documents():
    titles = ["关节镜下无源手术器械产品注册审查指导原则", "麻醉面罩产品注册审查指导原则"]
    body = titles[0] + ("甲" * 700) + "\n" + titles[1] + ("乙" * 700)
    chunks = fulltext.split_bodies(body, titles)
    assert set(chunks) == set(titles)
    assert chunks[titles[1]].startswith(titles[1])


def test_word_field_and_match():
    assert "I" in fulltext.clean_word_text("证明（ = 1 \\* ROMAN \\* MERGEFORMAT I）")
    pages = [{
        "key": fulltext.normalize_key("医疗器械出口销售证明管理规定"),
        "body_len": 20,
        "title": "医疗器械出口销售证明管理规定",
    }]
    found = fulltext.find_page(pages, "国家药监局关于发布医疗器械出口销售证明管理规定的公告")
    assert found["title"] == "医疗器械出口销售证明管理规定"
    edition = [{
        "title": "医疗器械生产质量管理规范(2025年版)",
        "body_len": 5000,
    }]
    assert fulltext.find_page(edition, "医疗器械生产质量管理规范", "2014年修订") is None
    assert fulltext.find_page(edition, "医疗器械生产质量管理规范", "自2025年11月起施行")["body_len"] == 5000
    editions = [
        {"title": "医疗器械生产质量管理规范（2014年版）", "body_len": 1000},
        {"title": "医疗器械生产质量管理规范(2025年版)", "body_len": 5000},
    ]
    assert fulltext.find_page(editions, "医疗器械生产质量管理规范", "第一条", "2025-11-06")["body_len"] == 5000
    assert fulltext.find_page(editions, "医疗器械生产质量管理规范", "第一条", "2014-12-01")["body_len"] == 1000
    notice = [{
        "title": "国家药监局关于进一步加强医疗器械注册人委托生产监督管理的公告（2024年第38号）",
        "body_len": 3000,
    }]
    assert fulltext.find_page(
        notice, "关于进一步加强医疗器械注册人委托生产监督管理", "", "2024-04-11"
    )["body_len"] == 3000
    older = [{"title": "中心静脉导管注册审查指导原则", "body_len": 4000}]
    revised = "中心静脉导管注册审查指导原则（2026年修订版）"
    assert fulltext.find_page(older, revised, "", "2026-07-15") is None
    numbered = "心磁信号采集设备注册审查指导原则（2026年第24号）"
    assert fulltext.find_page(
        [{"title": "心磁信号采集设备注册审查指导原则", "body_len": 8000}],
        numbered, "", "2026-07-21",
    )["body_len"] == 8000


def test_cmde_release_parser_skips_drafts():
    html = """
    <a href="/article/988">国家药监局器审中心关于发布有源植入式医疗器械电池注册审查指导原则的通告（2026年第27号）</a>
    <a href="/article/983">关于公开征求《植入式神经刺激器注册审查指导原则（征求意见稿）》意见的通知</a>
    <div><p>附件：有源植入式医疗器械电池注册审查指导原则（<a href="https://www.cmde.org.cn/directory/web/cmde/images/1.docx">下载</a>）</p>
    <p>发布时间: 2026-10-08</p></div>
    """
    listed = fulltext.parse_cmde_release_list(html)
    assert [item["link"] for item in listed] == ["https://www.ydcmdei.org.cn/article/988"]
    published, downloads = fulltext.parse_cmde_article(html, "https://www.ydcmdei.org.cn/article/987")
    assert published == "2026-10-08"
    assert downloads[0]["title"] == "有源植入式医疗器械电池注册审查指导原则"
    assert downloads[0]["url"].endswith("1.docx")
    revised = """
    <p>附件：人表皮生长因子受体2基因扩增检测试剂（荧光原位杂交法）注册审查指导原则（2026年修订版）（<a href="https://www.cmde.org.cn/directory/web/cmde/images/2.docx">下载</a>）</p>
    <p>发布时间: 2026-10-13</p>
    """
    _, revised_downloads = fulltext.parse_cmde_article(revised, "https://www.ydcmdei.org.cn/article/988")
    assert revised_downloads[0]["title"].endswith("（2026年修订版）")


def test_render_page_has_catalog_fields():
    page = fulltext.render_page(
        doc_id="nmpa-gov-test",
        title="麻醉面罩产品注册审查指导原则",
        section="nmpa/guidance",
        doc_number="2021年第102号",
        source_url="https://www.gov.cn/zhengce/content.htm",
        attachment="https://www.gov.cn/a.doc",
        published="2021-12-17",
        body="第一条 适用范围。",
    )
    data, body = fulltext.split_front_matter(page)
    assert data["category"] == "nmpa/guidance"
    assert data["document_number"] == "2021年第102号"
    assert "第一条" in body


def test_statute_is_not_filed_as_guidance():
    assert fulltext.target_section("医疗器械临床试验质量管理规范") == "nmpa/regulations"
    assert fulltext.target_section("医疗器械监督管理条例") == "nmpa/regulations"
    assert fulltext.target_section("脉搏血氧仪注册审查指导原则（2026年修订版）") == "nmpa/guidance"
    notice = "国家药监局关于实施《医疗器械临床试验质量管理规范》有关事项的通告"
    assert fulltext.is_statute_title(notice) is False
    assert fulltext.target_section(notice, "nmpa/guidance") == "nmpa/guidance"


def test_reflow_splits_paragraphs_and_builds_tables():
    prose = "第一条 适用范围。\n第二条 应当符合。\n"
    assert fulltext.reflow_markdown(prose) == "第一条 适用范围。\n\n第二条 应当符合。"
    loose = "危险分类 | 危险 | 伤害\n\n能量危险 | 电磁能 | 电击\n"
    table = fulltext.reflow_markdown(loose)
    assert table.startswith("| 危险分类 | 危险 | 伤害 |")
    assert "| --- | --- | --- |" in table
    assert "能量危险" in table


def test_docx_keeps_paragraphs_tables_and_images(tmp_path):
    from io import BytesIO
    from docx import Document
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
        "0000000c4944415408d763f8cfff000005fe02fe021c1a6a0000000049454e44ae426082"
    )
    image = tmp_path / "dot.png"
    image.write_bytes(png)
    doc = Document()
    doc.add_paragraph("第一条 适用范围。")
    doc.add_paragraph("第二条 应当符合。")
    doc.add_picture(str(image))
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "要求"
    table.rows[1].cells[0].text = "外观"
    table.rows[1].cells[1].text = "完整"
    buf = BytesIO()
    doc.save(buf)
    media = tmp_path / "media"
    text = fulltext.bytes_to_text(buf.getvalue(), media_dir=media, media_prefix="./media/sample")
    assert "第一条 适用范围。\n\n第二条 应当符合。" in text
    assert "| 项目 | 要求 |" in text
    assert "| --- | --- |" in text
    assert "| 外观 | 完整 |" in text
    assert "![](./media/sample/fig-1.png)" in text
    assert (media / "fig-1.png").is_file()
