#!/usr/bin/env python3
"""Turn a gov.cn policy page into catalog full text.

The weekly job can read www.gov.cn. Many 公告 and 办法 are the document
itself. 指导原则 are usually a .doc/.docx attachment on the same host,
which is not behind the NMPA challenge. This module extracts that text
and decides whether it should create, replace, or leave an existing page.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import urljoin, urlparse

try:
    import yaml
except ImportError:
    yaml = None

MIN_BODY = 600
MIN_KEEP = 1500
SECTIONS = ("nmpa/guidance", "nmpa/regulations")

_BOOK = re.compile(r"《([^》]{6,80})》")
_DOCNO = re.compile(
    r"((?:国药监[^\s《]{0,16})?〔\s*20\d{2}\s*〕\s*\d+\s*号|20\d{2}\s*年第\s*\d+\s*号)"
)
_WORD_FIELD = re.compile(
    r"=\s*\d+\s*\\\*\s*[A-Za-z]+\s*\\\*\s*MERGEFORMAT\s*"
)


def normalize_key(title: str) -> str:
    text = re.sub(r"\s+", "", title or "")
    return text.strip("《》")


def title_core(title: str) -> str:
    text = re.sub(r"[（(][^）)]*[）)]", "", normalize_key(title))
    return text


def _years(text: str) -> set[str]:
    return set(re.findall(r"20\d{2}", text or ""))


def clean_notice_title(title: str) -> str:
    text = re.sub(r"\s+", "", title or "")
    text = re.sub(r"^国家药监局", "", text)
    text = re.sub(r"^关于(发布|印发|修订发布|修订)", "", text)
    text = re.sub(r"的(公告|通告|通知|决定)(?=[（(]|$)", "", text)
    return text


def document_titles(notice_title: str) -> list[str]:
    books = []
    for name in _BOOK.findall(notice_title or ""):
        name = re.sub(r"\s+", "", name)
        if name and "附件" not in name and name not in books:
            books.append(name)
    if books:
        return books
    cleaned = clean_notice_title(notice_title)
    if len(cleaned) >= 6:
        return [cleaned]
    text = re.sub(r"\s+", "", notice_title or "")
    return [text] if text else []


def target_section(title: str, category_key: str = "") -> str:
    if re.search(r"指导原则|指南", title):
        return "nmpa/guidance"
    if re.search(r"办法|规定|条例|规章", title):
        return "nmpa/regulations"
    if category_key == "nmpa/guidance":
        return "nmpa/guidance"
    return "nmpa/regulations"


def document_number(text: str) -> str:
    match = _DOCNO.search(re.sub(r"\s+", "", text or ""))
    return re.sub(r"\s+", "", match.group(1)) if match else ""


def attachment_urls(html_text: str, page_url: str) -> list[str]:
    hrefs = re.findall(r'href=["\']([^"\']+)["\']', html_text or "", flags=re.I)
    found = []
    for href in hrefs:
        if not re.search(r"\.(?:doc|docx|pdf)(?:$|\?)", href, re.I):
            continue
        url = urljoin(page_url, href)
        if url not in found:
            found.append(url)
    return found


def parse_article(html_text: str, page_url: str) -> tuple[str, list[str]]:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return "", attachment_urls(html_text, page_url)
    soup = BeautifulSoup(html_text or "", "html.parser")
    node = soup.select_one("#UCAP-CONTENT") or soup.select_one(".pages_content")
    if node is None:
        return "", []
    text = clean_word_text(node.get_text("\n", strip=True))
    return text, attachment_urls(str(node), page_url)


def clean_word_text(text: str) -> str:
    cleaned = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    cleaned = _WORD_FIELD.sub("", cleaned)
    cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def needs_official_source(url: str) -> bool:
    host = (urlparse(url or "").hostname or "").lower()
    if not host:
        return True
    return not (host == "gov.cn" or host.endswith(".gov.cn") or host.endswith("cmde.org.cn"))


def plan_document(existing: dict | None, new_text: str) -> str:
    """create, replace, retarget, present, or no_fulltext."""
    new_len = len((new_text or "").strip())
    if existing is None:
        return "create" if new_len >= MIN_BODY else "no_fulltext"
    if existing.get("body_len", 0) >= MIN_KEEP:
        if needs_official_source(existing.get("source_url", "")):
            return "retarget"
        return "present"
    if new_len >= MIN_BODY and new_len > int(existing.get("body_len", 0)) + 200:
        return "replace"
    if needs_official_source(existing.get("source_url", "")) and new_len >= MIN_BODY:
        return "retarget"
    return "present" if existing.get("body_len", 0) else "no_fulltext"


def attachment_title(chunk: str) -> str:
    """Read the document name from the first lines of one 附件."""
    lines = [line.strip() for line in (chunk or "").splitlines() if line.strip()]
    title = ""
    for line in lines[:6]:
        if title and (line.startswith("本指导") or line.startswith("本指南") or line.startswith("本文件")):
            break
        title += line
        if re.search(r"(指导原则|指南|规定|办法|规范)$", title) and len(title) >= 6:
            return title[:80]
        if len(title) > 60:
            break
    return title[:80] if len(title) >= 6 else ""


def split_attachments(text: str) -> list[tuple[str, str]]:
    """Split a combined Word file that numbers each document as 附件1, 附件2."""
    matches = list(re.finditer(r"(?:^|\n)附件\s*[0-9０-９]+\s*(?:\n|$)", text or ""))
    if len(matches) < 2:
        return []
    pieces = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunk = text[match.end():end].strip()
        title = attachment_title(chunk)
        if title and len(chunk) >= MIN_BODY:
            pieces.append((title, chunk))
    return pieces if len(pieces) >= 2 else []


def split_bodies(text: str, titles: list[str]) -> dict[str, str]:
    """Slice one attachment when it contains several named documents."""
    body = text or ""
    if len(titles) <= 1:
        return {titles[0]: body} if titles else {}
    positions = []
    for title in titles:
        idx = body.find(title)
        if idx >= 0:
            positions.append((idx, title))
    positions.sort()
    if len(positions) < 2:
        return {}
    chunks = {}
    for i, (idx, title) in enumerate(positions):
        end = positions[i + 1][0] if i + 1 < len(positions) else len(body)
        chunk = body[idx:end].strip()
        if len(chunk) >= MIN_BODY:
            chunks[title] = chunk
    return chunks


def slug_for(title: str, published: str, prefix: str = "gov") -> str:
    year = (published or "")[:4] or "na"
    digest = hashlib.sha256(normalize_key(title).encode("utf-8")).hexdigest()[:10]
    return f"{prefix}-{year}-{digest}"


CMDE_RELEASE_URL = "https://www.ydcmdei.org.cn/article/category/release"


def clean_cmde_file_title(text: str) -> str:
    title = re.sub(r"\s+", "", text or "")
    title = re.sub(r"^附件[:：]?", "", title)
    title = re.sub(r"^[0-9]+[.、．]", "", title)
    title = re.sub(r"[（(]\s*下载\s*[）)]", "", title)
    return title.strip("：: ")


def parse_cmde_release_list(html_text: str) -> list[dict]:
    """Published guidance notices from the Yangtze CMDE mirror.

    The CMDE listing page itself returns an anti-bot challenge. This mirror
    lists the same 发布稿 and links each download to cmde.org.cn.
    """
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return []
    soup = BeautifulSoup(html_text or "", "html.parser")
    items = []
    seen = set()
    for anchor in soup.select("a"):
        href = anchor.get("href") or ""
        if not re.fullmatch(r"/article/\d+", href):
            continue
        title = re.sub(r"\s+", " ", anchor.get_text(" ", strip=True))
        if "征求" in title or ("通告" not in title and "指导原则" not in title):
            continue
        link = "https://www.ydcmdei.org.cn" + href
        if link in seen:
            continue
        seen.add(link)
        items.append({"title": title, "link": link})
    return items


def parse_cmde_article(html_text: str, page_url: str) -> tuple[str, list[dict]]:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return "", []
    soup = BeautifulSoup(html_text or "", "html.parser")
    match = re.search(r"发布时间[:：]\s*(20\d{2}-\d{2}-\d{2})", soup.get_text(" ", strip=True))
    published = match.group(1) if match else ""
    downloads = []
    for anchor in soup.select("a"):
        href = anchor.get("href") or ""
        if "cmde.org.cn" not in href or not re.search(r"\.(?:doc|docx|pdf)(?:$|\?)", href, re.I):
            continue
        parent = anchor.parent.get_text(" ", strip=True) if anchor.parent else ""
        title = clean_cmde_file_title(parent)
        if len(title) < 6:
            continue
        downloads.append({"title": title, "url": urljoin(page_url, href)})
    return published, downloads


def render_page(
    *,
    doc_id: str,
    title: str,
    section: str,
    doc_number: str,
    source_url: str,
    attachment: str,
    published: str,
    body: str,
) -> str:
    safe_title = title.replace('"', "'")
    lines = [
        "---",
        f"id: {doc_id}",
        "title:",
        f'  zh: "{safe_title}"',
        "  en: ''",
        "regulation: nmpa",
        f"category: {section}",
        "status: active",
        f"source_url: {source_url}",
        f"source_url_verified: '{published}'",
        "source_format: official",
        f"document_number: '{doc_number}'",
        f"effective_date: '{published}'",
        "contributor: official-sync",
    ]
    if attachment:
        lines.append(f"source_attachment: {attachment}")
    lines.append("---")
    lines.append("")
    lines.append(f"# {title}")
    lines.append("")
    lines.append(clean_word_text(body))
    lines.append("")
    return "\n".join(lines)


def set_official_source(text: str, url: str, published: str) -> str:
    if re.search(r"(?m)^source_url:", text):
        text = re.sub(r"(?m)^source_url:.*$", f"source_url: {url}", text, count=1)
    if re.search(r"(?m)^source_url_verified:", text):
        text = re.sub(
            r"(?m)^source_url_verified:.*$",
            f"source_url_verified: '{published}'",
            text,
            count=1,
        )
    if re.search(r"(?m)^source_url_status:", text):
        text = re.sub(r"(?m)^source_url_status:.*$", "source_url_status: official", text, count=1)
    return text


def split_front_matter(text: str) -> tuple[dict, str]:
    if yaml is None or not text.startswith("---"):
        return {}, text
    end = text.find("\n---", 3)
    if end < 0:
        return {}, text
    data = yaml.safe_load(text[3:end]) or {}
    if not isinstance(data, dict):
        return {}, text[end + 4:]
    return data, text[end + 4:]


def title_of(data: dict, fallback: str = "") -> str:
    title = data.get("title")
    if isinstance(title, dict):
        return str(title.get("zh") or fallback)
    if isinstance(title, str) and title:
        return title
    return fallback


def load_pages(repo: Path) -> list[dict]:
    pages = []
    for section in SECTIONS:
        folder = repo / section
        if not folder.is_dir():
            continue
        for path in sorted(folder.glob("*.zh.md")):
            text = path.read_text(encoding="utf-8")
            data, body = split_front_matter(text)
            title = title_of(data, path.stem)
            pages.append({
                "section": section,
                "path": path,
                "slug": path.name[: -len(".zh.md")],
                "title": title,
                "key": normalize_key(title),
                "body_len": len(body.strip()),
                "source_url": str(data.get("source_url") or ""),
                "doc_number": str(data.get("document_number") or data.get("doc_number") or ""),
                "doc_id": str(data.get("id") or ""),
            })
    return pages


def find_page(pages: list[dict], title: str, body: str = "", published: str = "") -> dict | None:
    """Match a catalog page without collapsing a different edition onto it."""
    wanted = title_core(clean_notice_title(title)) or title_core(title)
    if len(wanted) < 6:
        return None
    incoming_years = _years(title + "\n" + (body or "")[:1500] + "\n" + (published or ""))
    matches = []
    for page in pages:
        page_core = title_core(clean_notice_title(page.get("title") or ""))
        if page_core != wanted:
            continue
        page_years = _years(page.get("title") or "")
        if _years(title) and not page_years:
            continue
        if page_years and incoming_years and page_years.isdisjoint(incoming_years):
            continue
        if page_years and not incoming_years:
            continue
        matches.append(page)
    if not matches:
        return None
    return max(matches, key=lambda page: page["body_len"])


def looks_like_file(data: bytes) -> bool:
    if data.startswith(b"%PDF") or data.startswith(b"PK"):
        return True
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return True
    return False


def bytes_to_text(data: bytes) -> str:
    if not looks_like_file(data):
        return ""
    if data.startswith(b"%PDF"):
        return clean_word_text(_pdf_text(data))
    candidates = []
    docx = data if data.startswith(b"PK") else _libreoffice_docx(data)
    if docx.startswith(b"PK"):
        candidates.append(_docx_text(docx))
    candidates.append(_plain_text_convert(data))
    return clean_word_text(max(candidates, key=len)) if candidates else ""


def _pdf_text(data: bytes) -> str:
    try:
        import fitz
    except ImportError:
        return ""
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    from io import BytesIO
    try:
        doc = Document(BytesIO(data))
    except Exception:
        return ""
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            seen = set()
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in seen:
                    seen.add(value)
                    cells.append(value)
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _libreoffice_docx(data: bytes) -> bytes:
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if not binary:
        return b""
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(data)
        result = subprocess.run(
            [binary, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
            capture_output=True, text=True, timeout=90,
        )
        if result.returncode != 0:
            return b""
        produced = list(Path(tmp).glob("*.docx"))
        return produced[0].read_bytes() if produced else b""


def _plain_text_convert(data: bytes) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(data)
        out = Path(tmp) / "input.txt"
        if shutil.which("textutil"):
            result = subprocess.run(
                ["textutil", "-convert", "txt", "-output", str(out), str(src)],
                capture_output=True, text=True, timeout=60,
            )
            if result.returncode == 0 and out.exists():
                return out.read_text(encoding="utf-8", errors="replace")
        binary = shutil.which("soffice") or shutil.which("libreoffice")
        if not binary:
            return ""
        result = subprocess.run(
            [binary, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, str(src)],
            capture_output=True, text=True, timeout=90,
        )
        if result.returncode != 0:
            return ""
        produced = list(Path(tmp).glob("*.txt"))
        if not produced:
            return ""
        return produced[0].read_text(encoding="utf-8", errors="replace")
